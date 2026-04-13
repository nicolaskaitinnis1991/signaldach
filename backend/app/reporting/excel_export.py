"""Excel Export — Export trading data, signals, and price history to XLSX.

Supports multiple sheets in a single workbook:
- Price History
- Signal History
- Trade History (paper trading)
- Portfolio Summary
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, numbers
from openpyxl.utils import get_column_letter

from app.database import get_db

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Styling constants
# ---------------------------------------------------------------------------

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
DATE_FORMAT = "YYYY-MM-DD HH:MM:SS"
PRICE_FORMAT = '#,##0.00'
PCT_FORMAT = '0.00%'


def _style_header(ws) -> None:
    """Apply bold white-on-blue styling to the first row."""
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT


def _auto_column_widths(ws) -> None:
    """Set column widths based on content length."""
    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        adjusted = min(max_len + 4, 40)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted


def _workbook_to_bytes(wb: Workbook) -> bytes:
    """Save workbook to bytes buffer."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Price History
# ---------------------------------------------------------------------------

def export_price_history(symbol: str = "BTC/USDT", days: int = 30) -> bytes:
    """Export price history as XLSX with formatted columns.

    Returns XLSX file bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Price History"

    headers = ["Timestamp", "Symbol", "Exchange", "Price (USD)", "Volume 24h",
               "Market Cap", "Change % 24h"]
    ws.append(headers)
    _style_header(ws)

    cutoff = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()

    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT ts, symbol, exchange, price, volume_24h, market_cap, change_pct_24h
                   FROM price_ticks
                   WHERE symbol = ? AND ts >= ?
                   ORDER BY ts DESC""",
                [symbol, cutoff],
            )
            rows = cur.fetchall()

        for row in rows:
            r = dict(row) if not isinstance(row, dict) else row
            ws.append([
                r.get("ts", ""),
                r.get("symbol", symbol),
                r.get("exchange", ""),
                r.get("price", 0),
                r.get("volume_24h", 0),
                r.get("market_cap", 0),
                r.get("change_pct_24h", 0),
            ])

        # Apply number formatting to data rows
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=4).number_format = PRICE_FORMAT  # Price
            ws.cell(row=row_idx, column=5).number_format = PRICE_FORMAT  # Volume
            ws.cell(row=row_idx, column=6).number_format = PRICE_FORMAT  # Market Cap
            ws.cell(row=row_idx, column=7).number_format = '0.00'  # Change %

    except Exception as e:
        logger.warning("Could not fetch price data: %s", e)
        ws.append(["No price data available", "", "", "", "", "", ""])

    _auto_column_widths(ws)
    return _workbook_to_bytes(wb)


# ---------------------------------------------------------------------------
# Signal History
# ---------------------------------------------------------------------------

def export_signal_history(
    symbol: str = "BTC/USDT",
    signal_type: str = "composite",
    days: int = 30,
) -> bytes:
    """Export signal history as XLSX.

    Returns XLSX file bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Signal History"

    headers = ["Timestamp", "Symbol", "Signal Type", "Score", "Confidence", "Reasoning"]
    ws.append(headers)
    _style_header(ws)

    cutoff = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()

    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT ts, symbol, signal_type, score, confidence, reasoning
                   FROM signal_scores
                   WHERE symbol = ? AND signal_type = ? AND ts >= ?
                   ORDER BY ts DESC""",
                [symbol, signal_type, cutoff],
            )
            rows = cur.fetchall()

        for row in rows:
            r = dict(row) if not isinstance(row, dict) else row
            ws.append([
                r.get("ts", ""),
                r.get("symbol", symbol),
                r.get("signal_type", signal_type),
                r.get("score", 0),
                r.get("confidence", 0),
                r.get("reasoning", ""),
            ])

        # Format score and confidence columns
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=4).number_format = '0.00'
            ws.cell(row=row_idx, column=5).number_format = '0.00'

    except Exception as e:
        logger.warning("Could not fetch signal data: %s", e)
        ws.append(["No signal data available", "", "", "", "", ""])

    _auto_column_widths(ws)
    return _workbook_to_bytes(wb)


# ---------------------------------------------------------------------------
# Trade History (Paper Trading)
# ---------------------------------------------------------------------------

def export_trade_history(portfolio_id: str = "default") -> bytes:
    """Export paper trading history as XLSX.

    Returns XLSX file bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Trade History"

    headers = ["ID", "Symbol", "Side", "Quantity", "Entry Price", "Exit Price",
               "PnL", "PnL %", "Opened At", "Closed At", "Status"]
    ws.append(headers)
    _style_header(ws)

    try:
        with get_db() as conn:
            cur = conn.cursor()
            # paper_positions table from paper_db.py
            cur.execute(
                """SELECT id, symbol, side, quantity, entry_price, exit_price,
                          realized_pnl, unrealized_pnl_pct, opened_at, closed_at, status
                   FROM paper_positions
                   ORDER BY opened_at DESC"""
            )
            rows = cur.fetchall()

        for row in rows:
            r = dict(row) if not isinstance(row, dict) else row
            ws.append([
                r.get("id", ""),
                r.get("symbol", ""),
                r.get("side", ""),
                r.get("quantity", 0),
                r.get("entry_price", 0),
                r.get("exit_price"),
                r.get("realized_pnl", 0),
                r.get("unrealized_pnl_pct", 0),
                r.get("opened_at", ""),
                r.get("closed_at", ""),
                r.get("status", ""),
            ])

        # Format price / PnL columns
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=4).number_format = '0.0000'  # Quantity
            ws.cell(row=row_idx, column=5).number_format = PRICE_FORMAT  # Entry
            ws.cell(row=row_idx, column=6).number_format = PRICE_FORMAT  # Exit
            ws.cell(row=row_idx, column=7).number_format = PRICE_FORMAT  # PnL
            ws.cell(row=row_idx, column=8).number_format = '0.00'  # PnL %

    except Exception as e:
        logger.warning("Could not fetch trade data: %s", e)
        ws.append(["No trade data available", "", "", "", "", "", "", "", "", "", ""])

    _auto_column_widths(ws)
    return _workbook_to_bytes(wb)


# ---------------------------------------------------------------------------
# Full Report (Multi-Sheet)
# ---------------------------------------------------------------------------

def export_full_report(symbol: str = "BTC/USDT", days: int = 30) -> bytes:
    """Export a multi-sheet XLSX workbook with all data.

    Sheets:
    - Price History
    - Signal History
    - Trade History
    - Portfolio Summary

    Returns XLSX file bytes.
    """
    wb = Workbook()

    # --- Sheet 1: Price History ---
    ws_prices = wb.active
    ws_prices.title = "Price History"
    ws_prices.append(["Timestamp", "Symbol", "Exchange", "Price (USD)", "Volume 24h",
                       "Market Cap", "Change % 24h"])
    _style_header(ws_prices)

    cutoff = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()

    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT ts, symbol, exchange, price, volume_24h, market_cap, change_pct_24h
                   FROM price_ticks
                   WHERE symbol = ? AND ts >= ?
                   ORDER BY ts DESC""",
                [symbol, cutoff],
            )
            for row in cur.fetchall():
                r = dict(row) if not isinstance(row, dict) else row
                ws_prices.append([
                    r.get("ts", ""), r.get("symbol", symbol), r.get("exchange", ""),
                    r.get("price", 0), r.get("volume_24h", 0),
                    r.get("market_cap", 0), r.get("change_pct_24h", 0),
                ])
    except Exception as e:
        logger.warning("Price data unavailable for full report: %s", e)

    for row_idx in range(2, ws_prices.max_row + 1):
        ws_prices.cell(row=row_idx, column=4).number_format = PRICE_FORMAT
        ws_prices.cell(row=row_idx, column=5).number_format = PRICE_FORMAT
        ws_prices.cell(row=row_idx, column=6).number_format = PRICE_FORMAT
        ws_prices.cell(row=row_idx, column=7).number_format = '0.00'
    _auto_column_widths(ws_prices)

    # --- Sheet 2: Signal History ---
    ws_signals = wb.create_sheet("Signal History")
    ws_signals.append(["Timestamp", "Symbol", "Signal Type", "Score", "Confidence", "Reasoning"])
    _style_header(ws_signals)

    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT ts, symbol, signal_type, score, confidence, reasoning
                   FROM signal_scores
                   WHERE symbol = ? AND ts >= ?
                   ORDER BY ts DESC""",
                [symbol, cutoff],
            )
            for row in cur.fetchall():
                r = dict(row) if not isinstance(row, dict) else row
                ws_signals.append([
                    r.get("ts", ""), r.get("symbol", symbol),
                    r.get("signal_type", ""), r.get("score", 0),
                    r.get("confidence", 0), r.get("reasoning", ""),
                ])
    except Exception as e:
        logger.warning("Signal data unavailable for full report: %s", e)

    for row_idx in range(2, ws_signals.max_row + 1):
        ws_signals.cell(row=row_idx, column=4).number_format = '0.00'
        ws_signals.cell(row=row_idx, column=5).number_format = '0.00'
    _auto_column_widths(ws_signals)

    # --- Sheet 3: Trade History ---
    ws_trades = wb.create_sheet("Trade History")
    ws_trades.append(["ID", "Symbol", "Side", "Quantity", "Entry Price", "Exit Price",
                       "PnL", "PnL %", "Opened At", "Closed At", "Status"])
    _style_header(ws_trades)

    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT id, symbol, side, quantity, entry_price, exit_price,
                          realized_pnl, unrealized_pnl_pct, opened_at, closed_at, status
                   FROM paper_positions
                   ORDER BY opened_at DESC"""
            )
            for row in cur.fetchall():
                r = dict(row) if not isinstance(row, dict) else row
                ws_trades.append([
                    r.get("id", ""), r.get("symbol", ""), r.get("side", ""),
                    r.get("quantity", 0), r.get("entry_price", 0),
                    r.get("exit_price"), r.get("realized_pnl", 0),
                    r.get("unrealized_pnl_pct", 0), r.get("opened_at", ""),
                    r.get("closed_at", ""), r.get("status", ""),
                ])
    except Exception as e:
        logger.warning("Trade data unavailable for full report: %s", e)

    for row_idx in range(2, ws_trades.max_row + 1):
        ws_trades.cell(row=row_idx, column=5).number_format = PRICE_FORMAT
        ws_trades.cell(row=row_idx, column=6).number_format = PRICE_FORMAT
        ws_trades.cell(row=row_idx, column=7).number_format = PRICE_FORMAT
        ws_trades.cell(row=row_idx, column=8).number_format = '0.00'
    _auto_column_widths(ws_trades)

    # --- Sheet 4: Portfolio Summary ---
    ws_portfolio = wb.create_sheet("Portfolio Summary")
    ws_portfolio.append(["Symbol", "Side", "Quantity", "Entry Price", "Current Price",
                          "Unrealized PnL %", "Status"])
    _style_header(ws_portfolio)

    try:
        with get_db() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT symbol, side, quantity, entry_price, current_price,
                          unrealized_pnl_pct, status
                   FROM paper_positions
                   WHERE status = 'open'
                   ORDER BY symbol"""
            )
            for row in cur.fetchall():
                r = dict(row) if not isinstance(row, dict) else row
                ws_portfolio.append([
                    r.get("symbol", ""), r.get("side", ""),
                    r.get("quantity", 0), r.get("entry_price", 0),
                    r.get("current_price", 0), r.get("unrealized_pnl_pct", 0),
                    r.get("status", "open"),
                ])
    except Exception as e:
        logger.warning("Portfolio data unavailable for full report: %s", e)

    for row_idx in range(2, ws_portfolio.max_row + 1):
        ws_portfolio.cell(row=row_idx, column=4).number_format = PRICE_FORMAT
        ws_portfolio.cell(row=row_idx, column=5).number_format = PRICE_FORMAT
        ws_portfolio.cell(row=row_idx, column=6).number_format = '0.00'
    _auto_column_widths(ws_portfolio)

    return _workbook_to_bytes(wb)
